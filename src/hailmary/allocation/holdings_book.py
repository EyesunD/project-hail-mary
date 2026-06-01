"""Build the Stashaway portfolio book directly from ``data/holding.xlsx``.

Single-source ingestion that replaces the PDF-driven
``parse_statement → from_parsed`` path for the main analytic flow. Joins:

- ``PortfolioValue``: portfolio name, base currency, dated SGD + base-currency totals
- ``Holdings``: per-portfolio composition with dated ``Target %`` weights
- ``book_config``: role tags + per-portfolio Stashaway management fees (defaults)

The chosen ``Target %`` column is the latest ``YYYY-MM-DD`` snapshot whose date
is ``<= as_of``; the value columns ``MM/DD/YYYY (SGD|Base)`` must match
``as_of`` exactly. ``statement_fx_usd_sgd`` is derived from
``value_SGD / value_Base`` for USD-base portfolios — the rate Stashaway
actually applied, so no PDF lookup or Yahoo fetch is needed.
"""

from __future__ import annotations

import re
import warnings
from datetime import date, datetime
from pathlib import Path
from typing import Any

from hailmary.allocation.portfolios import Holding, Portfolio, Role
from hailmary.allocation.universe import STASHAWAY_UNIVERSE

_PORTFOLIO_ALIASES = {"SRS": "General SRS"}
_TARGET_COL_RE = re.compile(
    r"^\s*Target\s*%\s*(\d{4}-\d{2}-\d{2})?\s*$", re.IGNORECASE
)
_VALUE_COL_RE = re.compile(
    r"^\s*(\d{1,2}/\d{1,2}/\d{4})\s*\(\s*(SGD|Base)\s*\)\s*$", re.IGNORECASE
)
_HOLDING_LINK_RE = re.compile(r"/asset-details/([^/]+)/")
_CASH_ALIAS = {"USD": "CASH_USD", "SGD": "CASH_SGD"}

_VALUE_HEADER_ROW = 6
_HOLDINGS_HEADER_ROW = 7


class HoldingsBookError(ValueError):
    """Raised when ``holding.xlsx`` can't be parsed into a consistent book."""


def load_book(
    links_path: Path | str = Path("data/holding.xlsx"),
    *,
    as_of: date,
    roles_map: dict[str, set[Role]] | None = None,
    mgmt_fees: dict[str, float] | None = None,
) -> list[Portfolio]:
    """Build the portfolio book from ``holding.xlsx`` at ``as_of``.

    ``roles_map`` / ``mgmt_fees`` default to ``book_config.ROLES`` /
    ``MGMT_FEES_ANNUAL`` when ``None``. A portfolio without a roles entry
    defaults to ``{Role.HOLDING}``.

    Raises :class:`HoldingsBookError` if no value column matches ``as_of``
    or no ``Target %`` column is on/before ``as_of``.
    """
    try:
        import openpyxl
    except ImportError as exc:
        raise HoldingsBookError("openpyxl is required: pip install -e '.[allocation]'") from exc

    from hailmary.allocation.book_config import (
        MGMT_FEES_ANNUAL as DEFAULT_FEES,
        ROLES as DEFAULT_ROLES,
    )

    if roles_map is None:
        roles_map = DEFAULT_ROLES
    if mgmt_fees is None:
        mgmt_fees = DEFAULT_FEES

    path = Path(links_path)
    if not path.exists():
        raise HoldingsBookError(f"{path} does not exist")
    wb = openpyxl.load_workbook(path, data_only=False)
    for required in ("PortfolioValue", "Holdings"):
        if required not in wb.sheetnames:
            raise HoldingsBookError(f"Sheet {required!r} missing from {path}")

    pv = _load_portfolio_values(wb["PortfolioValue"], as_of)
    comp = _load_composition(wb["Holdings"], as_of)

    portfolios: list[Portfolio] = []
    for name in sorted(pv.keys()):
        currency, base_val, sgd_val = pv[name]
        rows = comp.get(name)
        if not rows:
            warnings.warn(
                f"No composition for {name!r} at {as_of}; skipping",
                stacklevel=2,
            )
            continue
        holdings: list[Holding] = []
        for ticker, weight in rows:
            md = STASHAWAY_UNIVERSE.get(ticker)
            if md is None:
                warnings.warn(
                    f"Unknown ticker {ticker!r} in {name!r}; dropping holding",
                    stacklevel=2,
                )
                continue
            holdings.append(
                Holding(
                    ticker=ticker,
                    weight=weight,
                    value=base_val * weight,
                    metadata=md,
                )
            )
        roles = set(roles_map.get(name, {Role.HOLDING}))
        meta: dict[str, Any] = {}
        fee = mgmt_fees.get(name)
        if fee is not None:
            meta["management_fee_annual"] = fee
        if currency == "USD" and base_val:
            meta["statement_fx_usd_sgd"] = sgd_val / base_val
        try:
            portfolios.append(
                Portfolio(
                    name=name,
                    statement_date=as_of,
                    total_value=base_val,
                    currency=currency,
                    holdings=holdings,
                    roles=roles,
                    metadata=meta,
                )
            )
        except ValueError as exc:
            warnings.warn(
                f"Portfolio {name!r} weight-sum check failed: {exc}; skipping",
                stacklevel=2,
            )
    return portfolios


def _load_portfolio_values(
    ws: Any, as_of: date
) -> dict[str, tuple[str, float, float]]:
    """Return ``{name: (currency, base_value, sgd_value)}`` for snapshot ``as_of``."""
    hdr_to_col: dict[str, int] = {}
    sgd_col: int | None = None
    base_col: int | None = None
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=_VALUE_HEADER_ROW, column=c).value
        if v is None:
            continue
        label = str(v).strip()
        hdr_to_col[label] = c
        m = _VALUE_COL_RE.match(label)
        if m:
            mm, dd, yyyy = m.group(1).split("/")
            col_date = date(int(yyyy), int(mm), int(dd))
            tag = m.group(2).upper()
            if col_date == as_of:
                if tag == "SGD":
                    sgd_col = c
                elif tag == "BASE":
                    base_col = c
    if sgd_col is None or base_col is None:
        available = sorted(
            {
                _VALUE_COL_RE.match(k).group(1)  # type: ignore[union-attr]
                for k in hdr_to_col
                if _VALUE_COL_RE.match(k)
            }
        )
        raise HoldingsBookError(
            f"PortfolioValue sheet has no SGD+Base columns for {as_of} "
            f"(available snapshots: {available})"
        )

    name_col = hdr_to_col.get("Public Market Investments")
    cur_col = hdr_to_col.get("Base Currency")
    if name_col is None or cur_col is None:
        raise HoldingsBookError(
            "PortfolioValue sheet must have 'Public Market Investments' and 'Base Currency' columns"
        )

    out: dict[str, tuple[str, float, float]] = {}
    for r in range(_VALUE_HEADER_ROW + 1, ws.max_row + 1):
        name = ws.cell(row=r, column=name_col).value
        if name is None or not str(name).strip():
            continue
        currency = ws.cell(row=r, column=cur_col).value
        sgd_v = ws.cell(row=r, column=sgd_col).value
        base_v = ws.cell(row=r, column=base_col).value
        if currency is None or sgd_v in (None, "") or base_v in (None, ""):
            continue
        try:
            out[str(name).strip()] = (
                str(currency).strip().upper(),
                float(base_v),
                float(sgd_v),
            )
        except (TypeError, ValueError):
            continue
    return out


def _load_composition(
    ws: Any, as_of: date
) -> dict[str, list[tuple[str, float]]]:
    """Return ``{portfolio_name: [(ticker, weight), ...]}`` from the latest
    ``Target % YYYY-MM-DD`` column whose effective date is ``<= as_of``.
    """
    hdr_to_col: dict[str, int] = {}
    target_cols: list[tuple[date, int]] = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=_HOLDINGS_HEADER_ROW, column=c).value
        if v is None:
            continue
        label = str(v).strip()
        hdr_to_col[label] = c
        m = _TARGET_COL_RE.match(label)
        if m and m.group(1):
            eff_date = datetime.strptime(m.group(1), "%Y-%m-%d").date()
            target_cols.append((eff_date, c))
    target_cols.sort(key=lambda t: t[0], reverse=True)
    chosen = next(((d, c) for d, c in target_cols if d <= as_of), None)
    if chosen is None:
        raise HoldingsBookError(
            f"Holdings sheet has no Target % column on/before {as_of} "
            f"(available: {[d.isoformat() for d, _ in target_cols]})"
        )
    eff_date, target_col = chosen

    port_col = hdr_to_col.get("Port")
    yt_col = hdr_to_col.get("Yahoo Ticker")
    link_col = hdr_to_col.get("Link")
    if port_col is None or yt_col is None:
        raise HoldingsBookError(
            "Holdings sheet must have 'Port' and 'Yahoo Ticker' columns"
        )

    agg: dict[str, dict[str, float]] = {}
    for r in range(_HOLDINGS_HEADER_ROW + 1, ws.max_row + 1):
        port = ws.cell(row=r, column=port_col).value
        weight = ws.cell(row=r, column=target_col).value
        if port is None or weight in (None, ""):
            continue
        try:
            w = float(weight)
        except (TypeError, ValueError):
            continue
        if w <= 0:
            continue
        name = _PORTFOLIO_ALIASES.get(str(port).strip(), str(port).strip())
        yt = ws.cell(row=r, column=yt_col).value
        ticker = str(yt).strip() if yt is not None and str(yt).strip() else ""
        if not ticker and link_col is not None:
            link_cell = ws.cell(row=r, column=link_col)
            url = link_cell.hyperlink.target if link_cell.hyperlink else ""
            m = _HOLDING_LINK_RE.search(str(url))
            if m:
                ticker = _CASH_ALIAS.get(m.group(1).upper(), "")
        if not ticker:
            continue
        bucket = agg.setdefault(name, {})
        bucket[ticker] = bucket.get(ticker, 0.0) + w
    return {name: list(weights.items()) for name, weights in agg.items()}
